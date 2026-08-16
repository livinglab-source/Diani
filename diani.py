# ==================================================
# BACKEND API V-FINAL (SUPPORT LIVE TEST & SKRIPSI)
# ==================================================
import os
import time
import csv
import glob
import shutil
from datetime import datetime
from collections import defaultdict
import cv2
import numpy as np
import torch
import torch.nn as nn
import torch.nn.functional as F
import uvicorn
from fastapi import FastAPI, Response, BackgroundTasks
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse

# ==========================
# SETUP APLIKASI
# ==========================
app = FastAPI(title="Dashboard Monitoring Pakcoy")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

OUTPUT_FOLDER = "hasil_deteksi_api"
DATA_FILE = "data_deteksi.csv"
MODEL_PATH_LOCAL = "unet_best_fold5.pth"
PIXELS_PER_GRAM = 1000

# 🔥 PENGATURAN WAKTU (CCTV HANYA AKTIF DI JAM INI)
SCAN_START_TIME_STR = "06:25:00"
SCAN_END_TIME_STR   = "06:40:00"

CAPTURE_INTERVAL_BIG = 120  # 120 detik = 2 menit
CROP_SIZE = 512
STRIDE = 400
DEVICE = "cuda" if torch.cuda.is_available() else "cpu"

os.makedirs(OUTPUT_FOLDER, exist_ok=True)
if not os.path.exists(DATA_FILE):
    with open(DATA_FILE, mode='w', newline='') as file:
        writer = csv.writer(file)
        writer.writerow(["Tanggal", "Jam", "PosX", "PosY", "PersenRusak", "HilangGram", "Filename"])

latest_image_path = None
session_history_files = []
is_running = False
model = None

# ==========================
# CCTV CONFIG
# ==========================
RTSP_URL = "rtsp://admin:admin@10.0.0.2:554/V_ENC_000"

# ==========================
# MODEL U-NET
# ==========================
class DoubleConv(nn.Module):
    def __init__(self, in_ch, out_ch):
        super().__init__()
        self.conv = nn.Sequential(
            nn.Conv2d(in_ch, out_ch, 3, padding=1), nn.BatchNorm2d(out_ch), nn.ReLU(inplace=True),
            nn.Conv2d(out_ch, out_ch, 3, padding=1), nn.BatchNorm2d(out_ch), nn.ReLU(inplace=True)
        )
    def forward(self, x): return self.conv(x)

class UNet(nn.Module):
    def __init__(self, n_class):
        super().__init__()
        self.d1 = DoubleConv(3, 64); self.d2 = DoubleConv(64, 128)
        self.d3 = DoubleConv(128, 256); self.d4 = DoubleConv(256, 512)
        self.bottleneck = DoubleConv(512, 1024)
        self.u4 = nn.ConvTranspose2d(1024, 512, 2, stride=2); self.uc4 = DoubleConv(1024, 512)
        self.u3 = nn.ConvTranspose2d(512, 256, 2, stride=2); self.uc3 = DoubleConv(512, 256)
        self.u2 = nn.ConvTranspose2d(256, 128, 2, stride=2); self.uc2 = DoubleConv(256, 128)
        self.u1 = nn.ConvTranspose2d(128, 64, 2, stride=2); self.uc1 = DoubleConv(128, 64)
        self.final = nn.Conv2d(64, n_class, 1)

    def forward(self, x):
        x1=self.d1(x); x2=self.d2(F.max_pool2d(x1, 2)); x3=self.d3(F.max_pool2d(x2, 2))
        x4=self.d4(F.max_pool2d(x3, 2)); xb=self.bottleneck(F.max_pool2d(x4, 2))
        x4=torch.cat([self.u4(xb), x4], 1); xu4=self.uc4(x4)
        x3=torch.cat([self.u3(xu4), x3], 1); xu3=self.uc3(x3)
        x2=torch.cat([self.u2(xu3), x2], 1); xu2=self.uc2(x2)
        x1=torch.cat([self.u1(xu2), x1], 1); xu1=self.uc1(x1)
        return self.final(xu1)

def load_model_unet():
    global model
    try:
        model_instance = UNet(3).to(DEVICE)
        if os.path.exists(MODEL_PATH_LOCAL):
            model_instance.load_state_dict(torch.load(MODEL_PATH_LOCAL, map_location=DEVICE))
            model_instance.eval()
            model = model_instance
    except Exception as e: pass

def predict_crop(crop_bgr):
    if model is None: return None, None
    img_rgb = cv2.cvtColor(crop_bgr, cv2.COLOR_BGR2RGB)
    img_tensor = torch.from_numpy(img_rgb).permute(2, 0, 1).float() / 255.0
    img_tensor = img_tensor.unsqueeze(0).to(DEVICE)
    with torch.no_grad():
        output = model(img_tensor)
        pred_mask = torch.argmax(output, dim=1).cpu().squeeze().numpy()
    return img_rgb, pred_mask

def save_statistics(timestamp_obj, x, y, pred_mask, filename):
    try:
        piksel_daun_sehat = np.sum(pred_mask == 1)
        piksel_daun_rusak = np.sum(pred_mask == 2)
        total_piksel_daun = piksel_daun_sehat + piksel_daun_rusak

        if total_piksel_daun > 0: damage_pct = round((piksel_daun_rusak / total_piksel_daun) * 100, 2)
        else: damage_pct = 0.0

        hilang_gram = round(piksel_daun_rusak / PIXELS_PER_GRAM, 2)
        tanggal = timestamp_obj.strftime('%Y-%m-%d')
        jam = timestamp_obj.strftime('%H:%M:%S')

        with open(DATA_FILE, mode='a', newline='') as file:
            writer = csv.writer(file)
            writer.writerow([tanggal, jam, x, y, damage_pct, hilang_gram, filename])
    except Exception as e: pass

# ==========================
# LOOP DETEKSI CCTV (DIUBAH AGAR VPS TIDAK MACET)
# ==========================
def detection_loop():
    global is_running, latest_image_path, session_history_files
    # Bikin seolah-olah dia butuh ngejepret langsung saat masuk jadwal pertama kali
    last_capture_time = time.time() - CAPTURE_INTERVAL_BIG

    try:
        while is_running:
            now = datetime.now().time()
            start_time = datetime.strptime(SCAN_START_TIME_STR, "%H:%M:%S").time()
            end_time = datetime.strptime(SCAN_END_TIME_STR, "%H:%M:%S").time()

            # Jika SEKARANG masuk dalam jadwal
            if start_time <= now <= end_time:
                current_time = time.time()
                # Jika SUDAH 2 MENIT sejak jepretan terakhir
                if current_time - last_capture_time >= CAPTURE_INTERVAL_BIG:

                    # Buka kamera, ambil 1 foto, lalu LANGSUNG TUTUP (hemat memori & CPU VPS)
                    cap = cv2.VideoCapture(RTSP_URL)
                    ret, frame = False, None
                    if cap.isOpened():
                        # PERBAIKAN 1: Kuras 5 frame pertama agar gambar benar-benar LIVE (Anti nyangkut)
                        for _ in range(5):
                            cap.read()
                        ret, frame = cap.read()
                        cap.release()

                    if ret and frame is not None:
                        last_capture_time = time.time()
                        now_obj = datetime.now()
                        timestamp_str = now_obj.strftime('%Y%m%d_%H%M%S')
                        today_folder = now_obj.strftime('%Y-%m-%d')

                        daily_path = os.path.join(OUTPUT_FOLDER, today_folder)
                        os.makedirs(daily_path, exist_ok=True)
                        big_frame = frame.copy()

                        h, w, _ = big_frame.shape
                        seen_crops = set() # PERBAIKAN 2: Satpam Koordinat Potongan

                        for y in range(0, h, STRIDE):
                            for x in range(0, w, STRIDE):
                                if not is_running: break
                                x1=x; y1=y; x2=x1+CROP_SIZE; y2=y1+CROP_SIZE
                                if x2 > w: x2 = w; x1 = w - CROP_SIZE
                                if y2 > h: y2 = h; y1 = h - CROP_SIZE

                                # PERBAIKAN 2: Jika koordinat sudah pernah dipotong di frame ini, LEWATI (Skip)
                                crop_coord = (x1, y1, x2, y2)
                                if crop_coord in seen_crops:
                                    continue
                                seen_crops.add(crop_coord)

                                crop_bgr = big_frame[y1:y2, x1:x2]
                                img_rgb_512, pred_mask = predict_crop(crop_bgr)

                                if img_rgb_512 is not None and pred_mask is not None:
                                    colors_bgr = np.array([[0,0,0], [0,255,0], [255,255,255]], dtype=np.uint8)
                                    mask_display = colors_bgr[pred_mask]
                                    img_display = cv2.cvtColor(img_rgb_512, cv2.COLOR_RGB2BGR)
                                    combined_result = np.hstack((img_display, mask_display))

                                    filename = f"scan_{timestamp_str}_pos_{x}_{y}_seg.jpg"
                                    filepath = os.path.join(daily_path, filename)
                                    if cv2.imwrite(filepath, combined_result):
                                        latest_image_path = filepath
                                        if filename not in session_history_files:
                                            session_history_files.append(filename)
                                        save_statistics(now_obj, x, y, pred_mask, filename)
                    else:
                        # Gagal menangkap gambar, istirahat 1 detik agar tidak error
                        time.sleep(1)
                else:
                    # Masih dalam rentang jadwal, tapi BELUM 2 menit.
                    # Istirahat 10 detik agar VPS sangat ringan.
                    time.sleep(10)
            else:
                # DI LUAR JADWAL. Mesin istirahat 10 detik agar adem.
                time.sleep(10)
    except Exception as e: pass
    finally:
        is_running = False

# ==========================
# ENDPOINT API
# ==========================
@app.on_event("startup")
async def startup_event():
    load_model_unet()
    global session_history_files
    today_str = datetime.now().strftime('%Y-%m-%d')
    today_folder = os.path.join(OUTPUT_FOLDER, today_str)

    if os.path.exists(today_folder):
        images = glob.glob(f"{today_folder}/*_seg.jpg")
        for img_path in images:
            filename = os.path.basename(img_path)
            try:
                jam_file = filename.split('_')[2]
                start_compact = SCAN_START_TIME_STR.replace(":", "")
                end_compact = SCAN_END_TIME_STR.replace(":", "")
                if start_compact <= jam_file <= end_compact:
                    if filename not in session_history_files:
                        session_history_files.append(filename)
            except: continue
        session_history_files.sort()

@app.post("/start-detection")
def start_detection(background_tasks: BackgroundTasks):
    global is_running
    if model is None: return {"status": "ERROR"}
    if is_running: return {"status": "INFO"}
    is_running = True
    background_tasks.add_task(detection_loop)
    return {"status": "SUCCESS"}

@app.get("/")
def home(): return FileResponse("index.html")

@app.get("/logo.jpg")
def get_logo():
    if os.path.exists("logo.jpg"): return FileResponse("logo.jpg", media_type="image/jpeg")
    return Response(status_code=404)

@app.get("/latest-image")
def get_latest_image_route():
    global latest_image_path
    if latest_image_path and os.path.exists(latest_image_path):
        headers = {'Cache-Control': 'no-cache'}
        return FileResponse(latest_image_path, media_type="image/jpeg", headers=headers)
    elif session_history_files:
        path = os.path.join(OUTPUT_FOLDER, datetime.now().strftime('%Y-%m-%d'), session_history_files[-1])
        if os.path.exists(path): return FileResponse(path, media_type="image/jpeg")
    return Response(status_code=404)

@app.get("/image/{filename}")
def get_specific_image(filename: str):
    for root, dirs, files in os.walk(OUTPUT_FOLDER):
        if filename in files:
            return FileResponse(os.path.join(root, filename), media_type="image/jpeg")
    return Response(status_code=404)

@app.get("/history")
def get_history_list():
    data = [{"filename": f, "pct": 0} for f in reversed(session_history_files)]
    return JSONResponse(content=data)

@app.get("/stats-today")
def get_stats_today():
    today_str = datetime.now().strftime('%Y-%m-%d')
    total_scans, total_damaged_scans, total_damage_pct, total_hilang_gram = 0, 0, 0, 0
    high_damage_areas = []

    if not os.path.exists(DATA_FILE):
        return JSONResponse(content={"summary": "Belum ada data hari ini.", "hotspots": []})

    with open(DATA_FILE, mode='r') as file:
        reader = csv.DictReader(file)
        for row in reader:
            if row["Tanggal"] != today_str: continue

            jam_str = row["Jam"]
            if not (SCAN_START_TIME_STR <= jam_str <= SCAN_END_TIME_STR):
                continue

            try:
                pct = float(row["PersenRusak"])
                gram = float(row.get("HilangGram", 0))
                total_scans += 1

                if pct > 0:
                    total_damaged_scans += 1
                    total_damage_pct += pct
                    total_hilang_gram += gram
                    high_damage_areas.append({
                        "x": row["PosX"], "y": row["PosY"], "pct": pct, "gram": gram, "filename": row["Filename"]
                    })
            except: continue

    avg_damage = round(total_damage_pct / total_damaged_scans, 2) if total_damaged_scans > 0 else 0
    total_hilang_gram = round(total_hilang_gram, 2)

    # PERMINTAAN 1: Sorting berdasarkan HilangGram tertinggi
    high_damage_areas.sort(key=lambda item: item['gram'], reverse=True)

    unique_hotspots = []
    seen_coords = set()
    for item in high_damage_areas:
        coord = (item['x'], item['y'])
        if coord not in seen_coords:
            unique_hotspots.append(item)
            seen_coords.add(coord)
        if len(unique_hotspots) == 4: break

    # LOGIKA TEKS STATUS DINAMIS IDE DARI DIANI
    now_time = datetime.now().time()
    start_time = datetime.strptime(SCAN_START_TIME_STR, "%H:%M:%S").time()
    end_time = datetime.strptime(SCAN_END_TIME_STR, "%H:%M:%S").time()

    if now_time < start_time:
        status_pemindaian = "⏳ Menunggu jadwal pemindaian"
    elif start_time <= now_time <= end_time:
        status_pemindaian = "🎥 CCTV sedang LIVE dan memindai"
    else:
        status_pemindaian = "✅ Pemindaian selesai"

    waktu_teks = f"{SCAN_START_TIME_STR[:5]} - {SCAN_END_TIME_STR[:5]}"
    summary = f"{status_pemindaian}. Total {total_scans} pemotongan gambar dari pukul {waktu_teks} telah diproses. Estimasi berat hilang: {total_hilang_gram} Gram."

    return JSONResponse(content={
        "date": today_str,
        "summary": summary,
        "average_damage_today": avg_damage,
        "total_hilang_gram": total_hilang_gram,
        "hotspots": unique_hotspots
    })

# PERMINTAAN 2: Endpoint Khusus Download Excel dengan Diagram Batang
@app.get("/download-excel")
def download_excel():
    try:
        import openpyxl
        from openpyxl.chart import BarChart, Reference
    except ImportError:
        return JSONResponse(status_code=500, content={"message": "Silakan install modul openpyxl dengan cara ketik 'pip install openpyxl' di terminal VPS."})

    wb = openpyxl.Workbook()
    ws_data = wb.active
    ws_data.title = "Rekap Harian"

    daily_stats = defaultdict(lambda: {'total_pct': 0.0, 'count': 0, 'total_gram': 0.0})

    if os.path.exists(DATA_FILE):
        with open(DATA_FILE, 'r') as f:
            reader = csv.DictReader(f)
            for row in reader:
                tanggal = row["Tanggal"]
                try:
                    pct = float(row["PersenRusak"])
                    gram = float(row.get("HilangGram", 0))

                    daily_stats[tanggal]['total_pct'] += pct
                    daily_stats[tanggal]['count'] += 1
                    daily_stats[tanggal]['total_gram'] += gram
                except: continue

    ws_data.append(["Tanggal", "Rata-rata Kerusakan (%)", "Total Hilang (Gram)"])

    row_idx = 2
    for tgl in sorted(daily_stats.keys()):
        st = daily_stats[tgl]
        avg_pct = round(st['total_pct'] / st['count'], 2) if st['count'] > 0 else 0
        tot_gram = round(st['total_gram'], 2)
        ws_data.append([tgl, avg_pct, tot_gram])
        row_idx += 1

    ws_chart = wb.create_sheet(title="Diagram Batang")
    chart = BarChart()
    chart.title = "Rekapitulasi Kerusakan Daun (Per Hari)"
    chart.x_axis.title = "Tanggal"
    chart.y_axis.title = "Nilai"

    if row_idx > 2:
        data_ref = Reference(ws_data, min_col=2, min_row=1, max_col=3, max_row=row_idx-1)
        cats_ref = Reference(ws_data, min_col=1, min_row=2, max_row=row_idx-1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)

    ws_chart.add_chart(chart, "B2")

    excel_path = "Laporan_Skripsi_Pakcoy.xlsx"
    wb.save(excel_path)

    return FileResponse(
        excel_path,
        filename="Laporan_Skripsi_Pakcoy.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

# PERMINTAAN BARU: Endpoint Download Gambar Sekaligus (.ZIP)
@app.get("/download-images")
def download_images():
    if not os.path.exists(OUTPUT_FOLDER) or not os.listdir(OUTPUT_FOLDER):
        return JSONResponse(status_code=404, content={"message": "Belum ada gambar yang bisa diunduh."})

    zip_filename = "kumpulan_gambar_deteksi"
    shutil.make_archive(zip_filename, 'zip', OUTPUT_FOLDER)

    return FileResponse(
        f"{zip_filename}.zip",
        media_type="application/zip",
        filename="Gambar_Skripsi_Pakcoy.zip"
    )

# PERMINTAAN BARU: Endpoint Cadangan Download CSV Mentah
@app.get("/download-csv")
def download_csv():
    if os.path.exists(DATA_FILE):
        return FileResponse(DATA_FILE, media_type="text/csv", filename="data_deteksi.csv")
    return Response(status_code=404)

if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=8004)